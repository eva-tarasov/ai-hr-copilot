"""Text-only PDF and DOCX extraction; no OCR or external link access."""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal
from zipfile import ZipFile

import pymupdf
from docx import Document

ParseStatus = Literal["ready", "unsupported", "parse_error"]
MAX_EXTRACTED_CHARS = 2_000_000
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_DOCX_MEMBERS = 1000


@dataclass(frozen=True)
class ParseResult:
    status: ParseStatus
    text: str | None = None
    error: str | None = None
    page_count: int = 0
    source_pages: list[dict[str, object]] = field(default_factory=list)


def parse_document(path: Path, extension: str) -> ParseResult:
    if path.stat().st_size == 0:
        return ParseResult(status="parse_error", error="File is empty")
    with path.open("rb") as source:
        signature = source.read(8)

    if extension == ".pdf":
        if signature.startswith(b"PK"):
            return ParseResult(status="unsupported", error="File content does not match PDF extension")
        if not signature.startswith(b"%PDF-"):
            return ParseResult(status="parse_error", error="Invalid PDF content")
        return parse_pdf(path)
    if extension == ".docx":
        if signature.startswith(b"%PDF-"):
            return ParseResult(status="unsupported", error="File content does not match DOCX extension")
        if not signature.startswith(b"PK"):
            return ParseResult(status="parse_error", error="Invalid DOCX content")
        return parse_docx(path)
    return ParseResult(status="unsupported", error="Only PDF and DOCX files are supported")


def parse_pdf(path: Path) -> ParseResult:
    pages: list[dict[str, object]] = []
    total_chars = 0
    try:
        with pymupdf.open(path) as document:
            page_count = len(document)
            for page_number, page in enumerate(document, start=1):
                page_text = page.get_text("text", sort=True)
                total_chars += len(page_text)
                if total_chars > MAX_EXTRACTED_CHARS:
                    return ParseResult(status="parse_error", error="Extracted text exceeds safe limit")
                pages.append({"page": page_number, "text": page_text})
    except Exception:
        return ParseResult(status="parse_error", error="Could not parse PDF")

    text = "\n\n".join(str(page["text"]) for page in pages).strip()
    if not text:
        return ParseResult(
            status="unsupported", error="No extractable PDF text; OCR is not available", page_count=page_count
        )
    return ParseResult(status="ready", text=text, page_count=page_count, source_pages=pages)


def parse_docx(path: Path) -> ParseResult:
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_DOCX_MEMBERS or sum(
                member.file_size for member in members
            ) > MAX_DOCX_UNCOMPRESSED_BYTES:
                return ParseResult(status="parse_error", error="DOCX exceeds safe parsing limits")
            names = {member.filename for member in members}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                return ParseResult(status="unsupported", error="File is not a DOCX document")

        document = Document(path)
        parts = [paragraph.text.strip() for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text.strip() for cell in row.cells)
        text = "\n".join(part for part in parts if part)
        if len(text) > MAX_EXTRACTED_CHARS:
            return ParseResult(status="parse_error", error="Extracted text exceeds safe limit")
    except Exception:
        return ParseResult(status="parse_error", error="Could not parse DOCX")

    if not text:
        return ParseResult(status="unsupported", error="No extractable DOCX text")
    return ParseResult(status="ready", text=text)
